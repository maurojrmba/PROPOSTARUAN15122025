import os
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database import listar_pendentes, COLUNAS


def gerar_excel() -> bytes:
    """Gera um arquivo Excel com todos os pagamentos pendentes, agrupados por obra."""
    registros = listar_pendentes()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pagamentos Pendentes"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="E65100")  # laranja escuro
    section_fill = PatternFill("solid", fgColor="FFF3E0")  # laranja claro
    thin = Side(style="thin", color="BDBDBD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    # Título
    ws.merge_cells("A1:J1")
    titulo = ws["A1"]
    titulo.value = f"RELATÓRIO DE PAGAMENTOS PENDENTES — {datetime.now().strftime('%d/%m/%Y')}"
    titulo.font = Font(bold=True, size=13, color="E65100")
    titulo.alignment = center
    ws.row_dimensions[1].height = 28

    # Cabeçalho
    colunas = COLUNAS  # ["ID","DATA","TIPO","VALOR","DESCRIÇÃO","OBRA","PIX","WHATSAPP","STATUS","IMAGEM"]
    ws.append([])  # linha 2 vazia
    ws.append(colunas)  # linha 3
    for col_idx, col_name in enumerate(colunas, start=1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # Agrupa por obra
    obras: dict[str, list] = {}
    for r in registros:
        obra = r[5]  # coluna 'obra'
        obras.setdefault(obra, []).append(r)

    row_num = 4
    total_geral = 0.0

    for obra, itens in obras.items():
        # Linha de seção (nome da obra)
        ws.merge_cells(f"A{row_num}:J{row_num}")
        cell = ws.cell(row=row_num, column=1)
        cell.value = f"  {obra}"
        cell.font = Font(bold=True, size=11, color="BF360C")
        cell.fill = section_fill
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row_num].height = 22
        row_num += 1

        subtotal = 0.0
        for item in itens:
            for col_idx, value in enumerate(item, start=1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            # Tenta somar valor
            try:
                valor_str = str(item[3]).replace("R$", "").replace(".", "").replace(",", ".").strip()
                subtotal += float(valor_str)
            except ValueError:
                pass
            row_num += 1

        # Subtotal da obra
        ws.merge_cells(f"A{row_num}:C{row_num}")
        ws.cell(row=row_num, column=1).value = f"SUBTOTAL — {obra}"
        ws.cell(row=row_num, column=1).font = Font(bold=True)
        ws.cell(row=row_num, column=4).value = f"R$ {subtotal:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        ws.cell(row=row_num, column=4).font = Font(bold=True)
        total_geral += subtotal
        row_num += 2

    # Total geral
    ws.merge_cells(f"A{row_num}:C{row_num}")
    ws.cell(row=row_num, column=1).value = "TOTAL GERAL"
    ws.cell(row=row_num, column=1).font = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(row=row_num, column=1).fill = PatternFill("solid", fgColor="E65100")
    ws.cell(row=row_num, column=4).value = f"R$ {total_geral:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    ws.cell(row=row_num, column=4).font = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(row=row_num, column=4).fill = PatternFill("solid", fgColor="E65100")

    # Largura das colunas
    larguras = [6, 16, 14, 12, 40, 22, 24, 16, 10, 20]
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def agendar_relatorio_sexta(app, admin_chat_id: int):
    """Agenda o envio automático do relatório toda sexta às 18h (horário de Brasília)."""
    from apscheduler.schedulers.asyncio import AsyncIOScheduler
    import asyncio

    scheduler = AsyncIOScheduler(timezone="America/Sao_Paulo")

    async def enviar():
        excel_bytes = gerar_excel()
        nome = f"relatorio_{datetime.now().strftime('%d%m%Y')}.xlsx"
        await app.bot.send_document(
            chat_id=admin_chat_id,
            document=excel_bytes,
            filename=nome,
            caption=f"📊 RELATÓRIO DE PAGAMENTOS PENDENTES — {datetime.now().strftime('%d/%m/%Y')}",
        )

    scheduler.add_job(enviar, "cron", day_of_week="fri", hour=18, minute=0)
    scheduler.start()
    return scheduler
